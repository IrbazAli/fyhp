"""
qwen_row_pipeline/storage.py
=============================================================================
Persistence, job directory structure, manifests, JSON export,
and genuine multi-sheet XLSX workbook generation via openpyxl.
=============================================================================
"""

import os
import json
import time
import uuid
import datetime
from typing import Dict, Any, List, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


JOBS_BASE_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "results", "qwen_row_pipeline")


def create_job_directory(source_filename: str) -> Dict[str, str]:
    """
    Creates an isolated job directory structure:
      results/qwen_row_pipeline/<job_id>/
        crops/
        context_crops/
        stages/
        exports/
    """
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    short_uuid = uuid.uuid4().hex[:6]
    job_id = f"job_{timestamp}_{short_uuid}"

    job_dir = os.path.join(JOBS_BASE_DIR, job_id)
    crops_dir = os.path.join(job_dir, "crops")
    context_dir = os.path.join(job_dir, "context_crops")
    stages_dir = os.path.join(job_dir, "stages")
    exports_dir = os.path.join(job_dir, "exports")

    os.makedirs(crops_dir, exist_ok=True)
    os.makedirs(context_dir, exist_ok=True)
    os.makedirs(stages_dir, exist_ok=True)
    os.makedirs(exports_dir, exist_ok=True)

    paths = {
        "job_id": job_id,
        "job_dir": job_dir,
        "crops_dir": crops_dir,
        "context_dir": context_dir,
        "stages_dir": stages_dir,
        "exports_dir": exports_dir,
        "manifest_path": os.path.join(job_dir, "manifest.json"),
        "canonical_image_path": os.path.join(job_dir, "canonical_page.png"),
        "localization_image_path": os.path.join(job_dir, "localization_image.png"),
        "transforms_path": os.path.join(job_dir, "transforms.json")
    }

    # Initialize manifest
    manifest = {
        "job_id": job_id,
        "created_at": datetime.datetime.now().isoformat(),
        "source_filename": source_filename,
        "model_info": {},
        "current_stage": "initialized",
        "stage_timings": {},
        "sections": [],
        "columns": [],
        "rows": [],
        "overall_flags": [],
        "stages": {
            "stage_a": {"status": "pending"},
            "stage_b": {"status": "pending", "refinements": []},
            "stage_c": {"status": "pending"},
            "stage_d": {"status": "pending"},
            "stage_e": {"status": "pending"},
            "stage_f": {"status": "pending"}
        }
    }
    save_manifest(paths["manifest_path"], manifest)
    return paths


def load_manifest(manifest_path: str) -> Dict[str, Any]:
    with open(manifest_path, "r", encoding="utf-8") as f:
        return json.load(f)


def save_manifest(manifest_path: str, manifest: Dict[str, Any]) -> None:
    temp_path = manifest_path + ".tmp"
    with open(temp_path, "w", encoding="utf-8") as f:
        json.dump(manifest, f, ensure_ascii=False, indent=2)
    os.replace(temp_path, manifest_path)


def save_stage_artifact(
    stages_dir: str,
    stage_name: str,
    raw_prompt: str,
    raw_response: str,
    generation_settings: Dict[str, Any],
    parsed_json: Optional[Dict[str, Any]] = None,
    duration_s: float = 0.0,
    is_truncated: bool = False
) -> None:
    """
    Saves raw prompt, raw response, settings, and parsed JSON for stage auditing.
    """
    req_file = os.path.join(stages_dir, f"{stage_name}_request.json")
    with open(req_file, "w", encoding="utf-8") as f:
        json.dump({
            "stage": stage_name,
            "timestamp": datetime.datetime.now().isoformat(),
            "duration_seconds": duration_s,
            "is_truncated": is_truncated,
            "generation_settings": generation_settings,
            "prompt": raw_prompt
        }, f, ensure_ascii=False, indent=2)

    raw_file = os.path.join(stages_dir, f"{stage_name}_raw_response.txt")
    with open(raw_file, "w", encoding="utf-8") as f:
        f.write(raw_response)

    if parsed_json is not None:
        parsed_file = os.path.join(stages_dir, f"{stage_name}_parsed.json")
        with open(parsed_file, "w", encoding="utf-8") as f:
            json.dump(parsed_json, f, ensure_ascii=False, indent=2)


def export_genuine_xlsx(
    manifest: Dict[str, Any],
    output_xlsx_path: str
) -> str:
    """
    Builds a genuine multi-sheet XLSX workbook with openpyxl.
    Sheets:
      1. Ledger: Reconstructed cells in page order with discovered columns, text formatting, RTL.
      2. Review: Flagged rows, ambiguous spans, conflict warnings.
      3. Provenance: Row IDs, crop paths, raw transcriptions, boundary revisions.
    All data values stored as explicit text strings to prevent Excel date/numeric corruption.
    """
    wb = Workbook()

    # Sheet 1: Ledger
    ws_ledger = wb.active
    ws_ledger.title = "Ledger"
    # Enable RTL reading order for Urdu ledger presentation
    ws_ledger.views.sheetView[0].rightToLeft = True

    # Styling palettes
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    accent_font = Font(name="Calibri", size=11, bold=True, color="00FFA3")
    data_font = Font(name="Calibri", size=10, color="000000")
    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC")
    )
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    right_align = Alignment(horizontal="right", vertical="center", wrap_text=True)

    # Title Banner
    ws_ledger.merge_cells("A1:G1")
    title_cell = ws_ledger["A1"]
    title_cell.value = f"Aasaan Khata — Ledger Extraction ({manifest.get('job_id')})"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="1E293B")
    ws_ledger.row_dimensions[1].height = 28

    ws_ledger.cell(row=2, column=1, value=f"Source Image: {manifest.get('source_filename')}").font = Font(size=9, italic=True)
    ws_ledger.cell(row=2, column=4, value=f"Exported: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}").font = Font(size=9, italic=True)

    # Build Header Row
    columns = manifest.get("columns", [])
    headers = ["Row ID", "Row Type"]
    for col in columns:
        headers.append(col.get("label", col.get("col_id", "Col")))
    headers.append("Raw Transcription")
    headers.append("Review Flags")

    row_idx = 4
    ws_ledger.row_dimensions[row_idx].height = 24
    for c_idx, h_text in enumerate(headers, start=1):
        c = ws_ledger.cell(row=row_idx, column=c_idx, value=h_text)
        c.fill = header_fill
        c.font = header_font
        c.alignment = center_align

    # Fill Ledger Data Rows
    rows = manifest.get("rows", [])
    for r in rows:
        row_idx += 1
        ws_ledger.row_dimensions[row_idx].height = 22

        r_id = str(r.get("row_id", ""))
        r_type = str(r.get("row_type", "transaction"))

        # Explicit string values
        c1 = ws_ledger.cell(row=row_idx, column=1, value=r_id)
        c1.alignment = center_align
        c1.font = data_font
        c1.border = thin_border

        c2 = ws_ledger.cell(row=row_idx, column=2, value=r_type)
        c2.alignment = center_align
        c2.font = data_font
        c2.border = thin_border

        # Columns
        stage_e = r.get("stage_e_fields") or {}
        cells = stage_e.get("cells", {}) if isinstance(stage_e, dict) else {}
        for c_idx, col in enumerate(columns, start=3):
            col_id = col.get("col_id", "")
            val = str(cells.get(col_id, ""))
            cell_obj = ws_ledger.cell(row=row_idx, column=c_idx)
            cell_obj.value = val
            cell_obj.data_type = 's'  # Explicit text to avoid Excel corruption
            cell_obj.alignment = right_align
            cell_obj.font = data_font
            cell_obj.border = thin_border

        # Raw transcription
        stage_d = r.get("stage_d_transcription") or {}
        raw_t = stage_d.get("raw_text", "") if isinstance(stage_d, dict) else ""
        c_raw = ws_ledger.cell(row=row_idx, column=len(columns) + 3, value=str(raw_t))
        c_raw.data_type = 's'
        c_raw.alignment = right_align
        c_raw.font = data_font
        c_raw.border = thin_border

        # Flags
        flags = list(r.get("validation_flags", []))
        if isinstance(stage_e, dict):
            flags.extend(stage_e.get("review_flags", []))
        c_flags = ws_ledger.cell(row=row_idx, column=len(columns) + 4, value=", ".join(flags))
        c_flags.alignment = center_align
        c_flags.font = Font(name="Calibri", size=9, color="DC2626" if flags else "4B5563")
        c_flags.border = thin_border

    # Auto-fit column widths
    for col in ws_ledger.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            max_len = max(max_len, len(val_str))
        ws_ledger.column_dimensions[col_letter].width = max(14, min(max_len + 4, 45))

    # Sheet 2: Review Sheet
    ws_review = wb.create_sheet(title="Review")
    ws_review.views.sheetView[0].rightToLeft = False
    rev_headers = ["Row ID", "Row Type", "Review Category", "Flags & Ambiguities", "Unassigned Text", "Current Bounds", "Crop Reference"]
    for c_idx, h_text in enumerate(rev_headers, start=1):
        c = ws_review.cell(row=1, column=c_idx, value=h_text)
        c.fill = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid")
        c.font = header_font
        c.alignment = center_align

    rev_row_idx = 1
    for r in rows:
        flags = list(r.get("validation_flags", []))
        stage_e = r.get("stage_e_fields") or {}
        ambiguous = []
        unassigned = ""
        if isinstance(stage_e, dict):
            flags.extend(stage_e.get("review_flags", []))
            ambiguous = stage_e.get("ambiguous_spans", [])
            unassigned = stage_e.get("unassigned_text", "")

        if flags or ambiguous or unassigned:
            rev_row_idx += 1
            ws_review.row_dimensions[rev_row_idx].height = 22
            ws_review.cell(row=rev_row_idx, column=1, value=str(r.get("row_id"))).alignment = center_align
            ws_review.cell(row=rev_row_idx, column=2, value=str(r.get("row_type"))).alignment = center_align
            ws_review.cell(row=rev_row_idx, column=3, value="Flagged").alignment = center_align
            ws_review.cell(row=rev_row_idx, column=4, value=", ".join(flags + [f"ambiguous: {x}" for x in ambiguous]))
            ws_review.cell(row=rev_row_idx, column=5, value=str(unassigned))
            bounds_str = f"top={r.get('current_bounds', {}).get('top')}, bot={r.get('current_bounds', {}).get('bottom')}"
            ws_review.cell(row=rev_row_idx, column=6, value=bounds_str).alignment = center_align
            ws_review.cell(row=rev_row_idx, column=7, value=str(r.get("crop_filename", "")))

    if rev_row_idx == 1:
        rev_row_idx = 2
        ws_review.cell(row=2, column=1, value="No review flags or ambiguities detected! All rows clean.").font = Font(italic=True, color="059669")

    for col in ws_review.columns:
        col_letter = get_column_letter(col[0].column)
        ws_review.column_dimensions[col_letter].width = 24

    # Sheet 3: Provenance Sheet
    ws_prov = wb.create_sheet(title="Provenance")
    ws_prov.views.sheetView[0].rightToLeft = False
    prov_headers = [
        "Row ID", "Crop File", "Proposed Bounds", "Current Bounds",
        "Manually Edited", "Stage B Refined", "Status", "Raw Transcription", "Structured JSON"
    ]
    for c_idx, h_text in enumerate(prov_headers, start=1):
        c = ws_prov.cell(row=1, column=c_idx, value=h_text)
        c.fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        c.font = header_font
        c.alignment = center_align

    for p_idx, r in enumerate(rows, start=2):
        ws_prov.row_dimensions[p_idx].height = 20
        prop_b = r.get("original_proposal", {})
        curr_b = r.get("current_bounds", {})
        prop_str = f"[{prop_b.get('top')}, {prop_b.get('bottom')}]"
        curr_str = f"[{curr_b.get('top')}, {curr_b.get('bottom')}]"

        stage_d = r.get("stage_d_transcription") or {}
        raw_t = stage_d.get("raw_text", "") if isinstance(stage_d, dict) else ""
        stage_e = r.get("stage_e_fields") or {}
        struct_json = json.dumps(stage_e.get("cells", {}), ensure_ascii=False) if isinstance(stage_e, dict) else ""

        ws_prov.cell(row=p_idx, column=1, value=str(r.get("row_id"))).alignment = center_align
        ws_prov.cell(row=p_idx, column=2, value=str(r.get("crop_filename")))
        ws_prov.cell(row=p_idx, column=3, value=prop_str).alignment = center_align
        ws_prov.cell(row=p_idx, column=4, value=curr_str).alignment = center_align
        ws_prov.cell(row=p_idx, column=5, value="Yes" if r.get("is_manually_edited") else "No").alignment = center_align
        ws_prov.cell(row=p_idx, column=6, value="Yes" if r.get("is_refined_stage_b") else "No").alignment = center_align
        ws_prov.cell(row=p_idx, column=7, value=str(r.get("status", "proposed"))).alignment = center_align
        ws_prov.cell(row=p_idx, column=8, value=str(raw_t))
        ws_prov.cell(row=p_idx, column=9, value=str(struct_json))

    for col in ws_prov.columns:
        col_letter = get_column_letter(col[0].column)
        ws_prov.column_dimensions[col_letter].width = 22

    wb.save(output_xlsx_path)
    return output_xlsx_path
