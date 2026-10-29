"""
tests/test_qwen_pipeline.py
=============================================================================
Unit and regression tests for Qwen Row Pipeline:
  - Factor-28 resolution scaling and coordinate inversions
  - Two-row context coordinate roundtrip
  - Row bounds validation, non-clamping out-of-bounds flagging, and overlap detection
  - Stable row identity and downstream stale-result invalidation
  - Robust JSON parsing and truncation handling
  - Genuine multi-sheet XLSX export with RTL and text-datatype preservation
=============================================================================
"""

import torch_patch
import os
import shutil
import tempfile
import unittest
from PIL import Image
from openpyxl import load_workbook

from qwen_row_pipeline.geometry import (
    compute_factor28_dimensions,
    prepare_localization_image,
    map_stage_a_coordinates,
    validate_row_bounds,
    create_two_row_context_crop,
    map_stage_b_coordinates,
    extract_row_crop
)
from qwen_row_pipeline.engine import parse_json_from_model_output
from qwen_row_pipeline.storage import export_genuine_xlsx


class TestQwenPipelineGeometry(unittest.TestCase):

    def test_factor28_dimensions(self):
        """Verify dimensions are strictly multiples of 28 and respect pixel caps."""
        test_cases = [
            (1920, 1080),
            (2400, 3200),
            (800, 600),
            (3000, 500),  # extreme wide
            (400, 3500)   # extreme tall
        ]
        for w, h in test_cases:
            target_w, target_h = compute_factor28_dimensions(w, h, max_pixels=602112)
            self.assertEqual(target_w % 28, 0, f"Width {target_w} not divisible by 28 for input ({w}, {h})")
            self.assertEqual(target_h % 28, 0, f"Height {target_h} not divisible by 28 for input ({w}, {h})")
            self.assertLessEqual(target_w * target_h, 602112, f"Total pixels exceeded max for ({w}, {h})")
            self.assertGreaterEqual(target_w, 28)
            self.assertGreaterEqual(target_h, 28)

    def test_stage_a_coordinate_mapping(self):
        """Verify coordinate scaling from model input back to canonical page."""
        transform = {
            "canonical_width": 2000,
            "canonical_height": 3000,
            "input_width": 500,
            "input_height": 750,
            "scale_x": 4.0,
            "scale_y": 4.0
        }
        raw_coords = {
            "rows": [
                {"top": 50.0, "bottom": 75.0, "ref_x": 250.0, "ref_y": 60.0, "row_type": "transaction"},
                {"top": 100.0, "bottom": 120.0, "ref_x": 200.0, "ref_y": 110.0, "row_type": "transaction"}
            ],
            "columns": [
                {"col_id": "c1", "label": "تفصیل", "guide_x": 150.0, "left": 100.0, "right": 300.0}
            ],
            "sections": [
                {"label": "12-03-2018", "top": 20.0, "bottom": 45.0}
            ]
        }
        mapped = map_stage_a_coordinates(raw_coords, transform)

        rows = mapped["rows"]
        self.assertEqual(len(rows), 2)
        # Check first row
        self.assertEqual(rows[0]["canonical_coords"]["top"], 200)
        self.assertEqual(rows[0]["canonical_coords"]["bottom"], 300)
        self.assertEqual(rows[0]["canonical_coords"]["left"], 0)
        self.assertEqual(rows[0]["canonical_coords"]["right"], 2000)
        self.assertEqual(rows[0]["canonical_coords"]["ref_x"], 1000)
        self.assertEqual(rows[0]["canonical_coords"]["ref_y"], 240)

        # Check column
        cols = mapped["columns"]
        self.assertEqual(cols[0]["guide_x_canonical"], 600)
        self.assertEqual(cols[0]["left_canonical"], 400)
        self.assertEqual(cols[0]["right_canonical"], 1200)

        # Check section
        secs = mapped["sections"]
        self.assertEqual(secs[0]["canonical_top"], 80)
        self.assertEqual(secs[0]["canonical_bottom"], 180)

    def test_hyphenated_and_string_coordinate_sanitization(self):
        """Verify robust handling of hyphenated strings ('281-'), trailing slashes, and non-numeric artifacts."""
        transform = {
            "canonical_width": 1200,
            "canonical_height": 1600,
            "input_width": 600,
            "input_height": 800,
            "scale_x": 2.0,
            "scale_y": 2.0
        }
        raw_coords = {
            "rows": [
                {"top": "271-", "bottom": "281-", "ref_x": "150-", "row_type": "transaction"},
                {"top": "50000/-", "bottom": "51000/-", "row_type": "transaction"},
                {"top": "N/", "bottom": "N/", "row_type": "unknown"}
            ]
        }
        mapped = map_stage_a_coordinates(raw_coords, transform)
        rows = mapped["rows"]
        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[0]["canonical_coords"]["top"], 542)
        self.assertEqual(rows[0]["canonical_coords"]["bottom"], 562)
        self.assertIn("sanitized_hyphenated_coordinate", rows[0]["raw_review_flags"])

    def test_bounds_validation_and_non_clamping(self):
        """
        Verify validation rules:
        - Inverted bounds flagged
        - Material out-of-bounds flagged (NOT silently clamped)
        - Overlaps and duplicates detected
        """
        img_w, img_h = 1000, 2000
        test_rows = [
            # 1. Normal valid row
            {"canonical_coords": {"top": 100, "bottom": 180}, "raw_review_flags": []},
            # 2. Overlapping row (with row 1)
            {"canonical_coords": {"top": 170, "bottom": 240}, "raw_review_flags": []},
            # 3. Inverted bounds
            {"canonical_coords": {"top": 400, "bottom": 350}, "raw_review_flags": []},
            # 4. Material out-of-bounds top
            {"canonical_coords": {"top": -50, "bottom": 80}, "raw_review_flags": []},
            # 5. Material out-of-bounds bottom
            {"canonical_coords": {"top": 1950, "bottom": 2150}, "raw_review_flags": []},
            # 6. Duplicate candidate (nearly 100% overlap with row 7)
            {"canonical_coords": {"top": 800, "bottom": 900}, "raw_review_flags": []},
            {"canonical_coords": {"top": 805, "bottom": 895}, "raw_review_flags": []}
        ]

        validated, issues = validate_row_bounds(test_rows, img_w, img_h)

        flags_by_idx = [r["validation_flags"] for r in validated]

        # Row 0 and Row 1 should have overlap flag
        self.assertTrue(any("overlap" in f for f in flags_by_idx[0]))
        self.assertTrue(any("overlap" in f for f in flags_by_idx[1]))

        # Row 2 (inverted)
        self.assertIn("inverted_or_empty_height", flags_by_idx[2])

        # Row 3 (material out-of-bounds top: -50px)
        self.assertIn("material_out_of_bounds_top", flags_by_idx[3])
        # Crucial check: must NOT have been silently clamped to 0!
        self.assertEqual(validated[3]["canonical_coords"]["top"], -50)

        # Row 4 (material out-of-bounds bottom: 2150px)
        self.assertIn("material_out_of_bounds_bottom", flags_by_idx[4])
        self.assertEqual(validated[4]["canonical_coords"]["bottom"], 2150)

        # Rows 5 & 6 (duplicate candidate)
        self.assertIn("duplicate_candidate", flags_by_idx[5])
        self.assertIn("duplicate_candidate", flags_by_idx[6])

    def test_two_row_context_transform_roundtrip(self):
        """
        Verify bidirectional coordinate mapping for Stage B refinement:
        model input -> context crop -> canonical full page.
        """
        tmp_dir = tempfile.mkdtemp()
        try:
            can_w, can_h = 1400, 2800
            canvas = Image.new("RGB", (can_w, can_h), color=(255, 255, 255))

            row_a_bounds = (500, 620)
            row_b_bounds = (610, 740)
            context_path = os.path.join(tmp_dir, "context_test.png")

            padding_y = 30
            ctx_trans = create_two_row_context_crop(
                canvas, row_a_bounds, row_b_bounds, context_path, padding_y=padding_y
            )

            ctx_box = ctx_trans["context_box_in_canonical"]
            expected_ctx_top = 500 - 30  # 470
            self.assertEqual(ctx_box["top"], expected_ctx_top)

            # Suppose Qwen predicts refined bounds in the model input coordinate space
            in_h = ctx_trans["model_input_dimensions"]["height"]
            scale_y = ctx_trans["scale_y"]

            # Synthetic prediction: row_a from top 10% to 45% of context crop
            pred_top_a_in = 0.10 * in_h
            pred_bot_a_in = 0.45 * in_h
            # row_b from 42% to 90%
            pred_top_b_in = 0.42 * in_h
            pred_bot_b_in = 0.90 * in_h

            ref_result = {
                "row_a": {"row_id": "row_001", "top": pred_top_a_in, "bottom": pred_bot_a_in, "ambiguity_flags": []},
                "row_b": {"row_id": "row_002", "top": pred_top_b_in, "bottom": pred_bot_b_in, "ambiguity_flags": ["touching_ink"]},
                "separation_notes": "Test notes",
                "conflict_detected": False
            }

            mapped = map_stage_b_coordinates(ref_result, ctx_trans)
            mapped_a = mapped["refined_rows"]["row_a"]["canonical_coords"]
            mapped_b = mapped["refined_rows"]["row_b"]["canonical_coords"]

            # Verify canonical coordinates are accurately mapped back
            expected_can_top_a = int(round(expected_ctx_top + (pred_top_a_in * scale_y)))
            expected_can_bot_a = int(round(expected_ctx_top + (pred_bot_a_in * scale_y)))
            self.assertEqual(mapped_a["top"], expected_can_top_a)
            self.assertEqual(mapped_a["bottom"], expected_can_bot_a)

            expected_can_top_b = int(round(expected_ctx_top + (pred_top_b_in * scale_y)))
            expected_can_bot_b = int(round(expected_ctx_top + (pred_bot_b_in * scale_y)))
            self.assertEqual(mapped_b["top"], expected_can_top_b)
            self.assertEqual(mapped_b["bottom"], expected_can_bot_b)
        finally:
            shutil.rmtree(tmp_dir)


class TestQwenPipelineStorageAndExcel(unittest.TestCase):

    def test_json_parsing_and_truncation_handling(self):
        """Verify robust parsing of JSON fences, trailing commas, and malformed strings."""
        # 1. Clean markdown fenced JSON
        s1 = '```json\n{"rows": [{"top": 10, "bottom": 20}]}\n```'
        p1 = parse_json_from_model_output(s1)
        self.assertIsNotNone(p1)
        self.assertEqual(len(p1["rows"]), 1)

        # 2. Raw JSON with surrounding conversation text
        s2 = 'Here is the extracted layout JSON:\n{"rows": [{"top": 30, "bottom": 40}]}\nHope this helps!'
        p2 = parse_json_from_model_output(s2)
        self.assertIsNotNone(p2)
        self.assertEqual(p2["rows"][0]["top"], 30)

        # 3. JSON with trailing comma
        s3 = '{"rows": [{"top": 50, "bottom": 60},],}'
        p3 = parse_json_from_model_output(s3)
        self.assertIsNotNone(p3)

        # 4. Truncated array recovery (e.g. last item cut off mid-string)
        s4 = '```json\n[\n  {"top": 10, "bottom": 20},\n  {"top": 20, "bottom": 30},\n  {"top": 30, "bottom": 4'
        p4 = parse_json_from_model_output(s4)
        self.assertIsNotNone(p4)
        self.assertEqual(len(p4["rows"]), 2)
        self.assertTrue(p4.get("is_partial_recovery"))

        # 5. Completely empty or unparseable with zero closed objects returns None
        s5 = '["top": 50'
        p5 = parse_json_from_model_output(s5)
        self.assertIsNone(p5)

        # 6. Model wrapped entire root object in an array: [{"coordinate_space": ..., "rows": [...]}]
        s6 = '```json\n[\n  {"coordinate_space": {"width": 672, "height": 896}, "rows": [{"top": 100, "bottom": 140}, {"top": 140, "bottom": 180}]}\n]\n```'
        p6 = parse_json_from_model_output(s6)
        self.assertIsNotNone(p6)
        self.assertEqual(len(p6.get("rows", [])), 2)
        self.assertEqual(p6["rows"][0]["top"], 100)

        # 7. Truncated nested structure cut off mid-row inside array-wrapped object
        s7 = '```json\n[\n  {"section_space": {"width": 672, "height": 896}, "rows": [\n    {"top": 100, "bottom": 140},\n    {"top": 140, "bottom": 180},\n    {"top": 180'
        p7 = parse_json_from_model_output(s7)
        self.assertIsNotNone(p7)
        self.assertEqual(len(p7.get("rows", [])), 2)
        self.assertTrue(p7.get("is_partial_recovery"))


    def test_genuine_xlsx_export(self):
        """
        Verify multi-sheet XLSX export:
        - Ledger, Review, and Provenance sheets exist
        - Urdu text preserved intact
        - Literal // preserved
        - Leading zeros preserved ('0042') without truncation to integer 42
        - Right-to-Left (RTL) enabled on Ledger sheet
        """
        tmp_dir = tempfile.mkdtemp()
        try:
            xlsx_path = os.path.join(tmp_dir, "test_export.xlsx")
            sample_manifest = {
                "job_id": "job_unit_test",
                "source_filename": "sample_ledger.jpg",
                "columns": [
                    {"col_id": "col_01", "label": "تاریخ (Date)"},
                    {"col_id": "col_02", "label": "تفصیل آمدن (Description)"},
                    {"col_id": "col_03", "label": "صفحہ (Folio)"},
                    {"col_id": "col_04", "label": "رقم روپیہ (Amount)"}
                ],
                "rows": [
                    {
                        "row_id": "row_001",
                        "row_type": "transaction",
                        "crop_filename": "row_001.png",
                        "original_proposal": {"top": 100, "bottom": 160},
                        "current_bounds": {"top": 100, "bottom": 160},
                        "is_manually_edited": False,
                        "is_refined_stage_b": False,
                        "validation_flags": [],
                        "stage_d_transcription": {
                            "raw_text": "11-03-018 کیش // 4024/-",
                            "cleaned_text": "11-03-018 کیش // 4024/-"
                        },
                        "stage_e_fields": {
                            "cells": {
                                "col_01": "11-03-018",
                                "col_02": "کیش",
                                "col_03": "007",  # leading zeros test!
                                "col_04": "4024/-"
                            },
                            "unassigned_text": "",
                            "ambiguous_spans": [],
                            "review_flags": []
                        }
                    },
                    {
                        "row_id": "row_002",
                        "row_type": "transaction",
                        "crop_filename": "row_002.png",
                        "original_proposal": {"top": 160, "bottom": 220},
                        "current_bounds": {"top": 165, "bottom": 225},
                        "is_manually_edited": True,
                        "is_refined_stage_b": False,
                        "validation_flags": ["overlap_with_row_001"],
                        "stage_d_transcription": {
                            "raw_text": "11-03-018 مونگ پھلی 40 11745/-",
                            "cleaned_text": "11-03-018 مونگ پھلی 40 11745/-"
                        },
                        "stage_e_fields": {
                            "cells": {
                                "col_01": "11-03-018",
                                "col_02": "مونگ پھلی //",  # literal ditto marks!
                                "col_03": "40",
                                "col_04": "11745/-"
                            },
                            "unassigned_text": "",
                            "ambiguous_spans": ["ambiguous_zero"],
                            "review_flags": ["cut_uncertainty"]
                        }
                    }
                ]
            }

            export_genuine_xlsx(sample_manifest, xlsx_path)
            self.assertTrue(os.path.exists(xlsx_path))

            # Inspect workbook
            wb = load_workbook(xlsx_path)
            sheet_names = wb.sheetnames
            self.assertIn("Ledger", sheet_names)
            self.assertIn("Review", sheet_names)
            self.assertIn("Provenance", sheet_names)

            # Check Ledger Sheet
            ws_ledger = wb["Ledger"]
            self.assertTrue(ws_ledger.views.sheetView[0].rightToLeft, "Ledger sheet must have RTL enabled")

            # Check Row 1 (row_idx 5 in sheet due to headers)
            folio_cell = ws_ledger.cell(row=5, column=5)  # Col 1: ID, Col 2: Type, Col 3: Date, Col 4: Desc, Col 5: Folio
            self.assertEqual(folio_cell.value, "007", "Leading zero must be preserved as string '007'")

            desc_cell_row2 = ws_ledger.cell(row=6, column=4)
            self.assertEqual(desc_cell_row2.value, "مونگ پھلی //", "Urdu Nastaliq and literal // must be preserved")

            # Check Review Sheet
            ws_review = wb["Review"]
            # Row 2 should list row_002 with flags and ambiguous spans
            self.assertIn("overlap_with_row_001", str(ws_review.cell(row=2, column=4).value))

            # Check Provenance Sheet
            ws_prov = wb["Provenance"]
            # Row 002 is manually edited
            self.assertEqual(ws_prov.cell(row=3, column=5).value, "Yes")
        finally:
            shutil.rmtree(tmp_dir)

    def test_row_identity_and_crop_edit_invalidation(self):
        """
        Verify that manual crop edits:
        1. Preserve stable row ID and original proposal
        2. Update current bounds and mark is_manually_edited=True
        3. Invalidate downstream Stage D and Stage E results as stale
        4. Regenerate the physical crop image
        5. Support clean restoration of original proposed bounds
        """
        import json
        from qwen_row_pipeline.engine import QwenRowPipelineRunner
        from qwen_row_pipeline.storage import JOBS_BASE_DIR, save_manifest

        test_job_id = "test_job_edit_invalidation"
        job_dir = os.path.join(JOBS_BASE_DIR, test_job_id)
        crops_dir = os.path.join(job_dir, "crops")
        os.makedirs(crops_dir, exist_ok=True)

        try:
            # Create synthetic canonical image 1000x1000
            canvas = Image.new("RGB", (1000, 1000), color=(240, 240, 240))
            can_path = os.path.join(job_dir, "canonical_page.png")
            canvas.save(can_path)

            # Extract initial crop
            crop_path = os.path.join(crops_dir, "row_001.png")
            init_crop = canvas.crop((0, 100, 1000, 200))
            init_crop.save(crop_path)

            manifest_path = os.path.join(job_dir, "manifest.json")
            manifest = {
                "job_id": test_job_id,
                "rows": [
                    {
                        "row_id": "row_001",
                        "row_type": "transaction",
                        "original_proposal": {"top": 100, "bottom": 200, "left": 0, "right": 1000},
                        "current_bounds": {"top": 100, "bottom": 200, "left": 0, "right": 1000},
                        "crop_filename": "row_001.png",
                        "crop_width": 1000,
                        "crop_height": 100,
                        "status": "proposed",
                        "is_manually_edited": False,
                        "is_refined_stage_b": False,
                        "validation_flags": [],
                        "stage_d_transcription": {
                            "raw_text": "Sample text",
                            "stale": False
                        },
                        "stage_e_fields": {
                            "cells": {"col_01": "Sample"},
                            "stale": False
                        }
                    }
                ]
            }
            save_manifest(manifest_path, manifest)

            runner = QwenRowPipelineRunner()

            # Execute crop edit: top 100 -> 120, bottom 200 -> 250
            res = runner.update_crop_boundary(test_job_id, "row_001", new_top=120, new_bottom=250)
            row_updated = res["row"]

            # 1. Original proposal MUST remain untouched
            self.assertEqual(row_updated["original_proposal"]["top"], 100)
            self.assertEqual(row_updated["original_proposal"]["bottom"], 200)

            # 2. Current bounds updated
            self.assertEqual(row_updated["current_bounds"]["top"], 120)
            self.assertEqual(row_updated["current_bounds"]["bottom"], 250)
            self.assertTrue(row_updated["is_manually_edited"])

            # 3. Downstream results invalidated
            self.assertTrue(row_updated["stage_d_transcription"]["stale"])
            self.assertTrue(row_updated["stage_e_fields"]["stale"])

            # 4. Physical crop resized to 250 - 120 = 130px
            with Image.open(crop_path) as reloaded_crop:
                self.assertEqual(reloaded_crop.size, (1000, 130))

            # 5. Restore original proposal
            res_restore = runner.restore_original_row_proposal(test_job_id, "row_001")
            row_restored = res_restore["row"]
            self.assertEqual(row_restored["current_bounds"]["top"], 100)
            self.assertEqual(row_restored["current_bounds"]["bottom"], 200)
            self.assertFalse(row_restored["is_manually_edited"])
            with Image.open(crop_path) as restored_crop:
                self.assertEqual(restored_crop.size, (1000, 100))

        finally:
            if os.path.exists(job_dir):
                shutil.rmtree(job_dir)

    def test_vision_preprocessor_patch_extraction(self):
        """
        Verify Vision Transformer patch extraction:
        - Exact patch dimensions: base 14x14, merged 28x28 (merge_size=2)
        - In-memory processing with base64 Data URL
        - Zero patch files written to disk
        - Downstream transcription commented down
        """
        from qwen_row_pipeline.engine import QwenRowPipelineRunner

        runner = QwenRowPipelineRunner()
        sample_path = "dataset/1.jpeg"
        if not os.path.exists(sample_path):
            self.skipTest("Sample dataset/1.jpeg not present for patch test")

        res = runner.extract_vision_patches(sample_path, "1.jpeg")
        self.assertEqual(res["status"], "success")
        self.assertEqual(res["mode"], "vision_patches")

        # Verify patch dimensions
        patch_dims = res["patch_dimensions"]
        self.assertEqual(patch_dims["base_patch_size"], [14, 14])
        self.assertEqual(patch_dims["merge_size"], 2)
        self.assertEqual(patch_dims["merged_token_size"], [28, 28])

        # Verify grid geometry
        grid = res["grid"]
        base_grid = grid["base_grid"]
        token_grid = grid["token_grid"]
        self.assertEqual(base_grid["total_patches"], base_grid["rows"] * base_grid["cols"])
        self.assertEqual(token_grid["total_tokens"], token_grid["rows"] * token_grid["cols"])
        self.assertEqual(token_grid["rows"], base_grid["rows"] // 2)
        self.assertEqual(token_grid["cols"], base_grid["cols"] // 2)

        # Verify in-memory base64 image (zero patch files written)
        self.assertTrue(res["image_data_url"].startswith("data:image/jpeg;base64,"))
        self.assertGreater(len(res["image_data_url"]), 1000)

        # Verify downstream translation is commented down
        res_transcribe = runner.transcribe_and_assign_reviewed_crops("dummy_job_id")
        self.assertEqual(res_transcribe["status"], "commented_down")


class TestKrakenBLLAPreprocessor(unittest.TestCase):

    def test_kraken_segmentation_structure(self):
        """Verify Kraken BLLA extracts baselines, polygons, bboxes and in-memory crops."""
        from qwen_row_pipeline.kraken_engine import KrakenBLLAPreprocessor
        sample_path = "dataset/1.jpeg"
        if not os.path.exists(sample_path):
            self.skipTest("Sample dataset/1.jpeg not present for Kraken test")

        preprocessor = KrakenBLLAPreprocessor()
        res = preprocessor.segment_image(sample_path, text_direction="horizontal-rl", source_filename="1.jpeg")

        self.assertEqual(res["status"], "success")
        self.assertEqual(res["mode"], "kraken_blla")
        self.assertEqual(res["model_name"], "blla.mlmodel")
        self.assertGreater(res["line_count"], 0)
        self.assertTrue(res["full_image_data_url"].startswith("data:image/jpeg;base64,"))

        first_line = res["lines"][0]
        self.assertIn("baseline", first_line)
        self.assertIn("boundary", first_line)
        self.assertIn("bbox", first_line)
        self.assertIn("crop_data_url", first_line)
        self.assertTrue(first_line["crop_data_url"].startswith("data:image/jpeg;base64,"))
        self.assertGreaterEqual(first_line["baseline_points_count"], 2)
        self.assertGreaterEqual(first_line["boundary_points_count"], 3)


class TestSegmentationBenchmark(unittest.TestCase):
    """Verifies all 5 segmentation candidate engines and the comparative benchmark runner."""

    def test_opencv_ruling_detector(self):
        """Verify OpenCV morphological ruling detector slices rows between ruled lines."""
        from segmentation_benchmark.opencv_ruling_engine import OpenCVRulingPreprocessor
        sample_path = "dataset/1.jpeg"
        if not os.path.exists(sample_path):
            self.skipTest("Sample dataset/1.jpeg not present")

        engine = OpenCVRulingPreprocessor()
        res = engine.segment_image(sample_path)
        self.assertEqual(res["status"], "success")
        self.assertGreater(res["rule_count"], 0)
        self.assertGreater(len(res["lines"]), 0)
        self.assertTrue(res["lines"][0]["crop_data_url"].startswith("data:image/jpeg;base64,"))
        self.assertIn("rulings", res)

    def test_table_transformer_detector(self):
        """Verify Table Transformer TATR extracts rows, columns, and table bounds."""
        from segmentation_benchmark.table_transformer_engine import TableTransformerPreprocessor
        sample_path = "dataset/1.jpeg"
        if not os.path.exists(sample_path):
            self.skipTest("Sample dataset/1.jpeg not present")

        engine = TableTransformerPreprocessor()
        res = engine.segment_image(sample_path)
        self.assertEqual(res["status"], "success")
        self.assertGreater(len(res["lines"]), 0)
        categories = {item.get("category", "").lower().replace("_", " ") for item in res["lines"]}
        self.assertTrue("table row" in categories or "table" in categories)

    def test_yolo_layout_detector(self):
        """Verify YOLO layout detector extracts document layout bounding boxes."""
        from segmentation_benchmark.yolo_engine import YOLOLayoutPreprocessor
        sample_path = "dataset/1.jpeg"
        if not os.path.exists(sample_path):
            self.skipTest("Sample dataset/1.jpeg not present")

        engine = YOLOLayoutPreprocessor()
        res = engine.segment_image(sample_path)
        self.assertEqual(res["status"], "success")
        self.assertIn("lines", res)

    def test_paddleocr_dbnet_detector(self):
        """Verify PaddleOCR DBNet text region detector extracts text polygons."""
        from segmentation_benchmark.paddle_det_engine import PaddleDBNetPreprocessor
        sample_path = "dataset/1.jpeg"
        if not os.path.exists(sample_path):
            self.skipTest("Sample dataset/1.jpeg not present")

        engine = PaddleDBNetPreprocessor()
        res = engine.segment_image(sample_path)
        self.assertEqual(res["status"], "success")
        self.assertGreater(len(res["lines"]), 0)
        self.assertIn("polygon", res["lines"][0])

    def test_comparative_benchmark_runner(self):
        """Verify unified benchmark runner produces a comparative scorecard across engines."""
        from segmentation_benchmark.benchmark_runner import run_comparative_benchmark
        sample_path = "dataset/1.jpeg"
        if not os.path.exists(sample_path):
            self.skipTest("Sample dataset/1.jpeg not present")

        res = run_comparative_benchmark(sample_path)
        self.assertEqual(res["status"], "success")
        self.assertIn("scorecard", res)
        self.assertIn("results", res)
        self.assertEqual(len(res["scorecard"]), 5)
        engine_keys = [sc["engine_key"] for sc in res["scorecard"]]
        self.assertIn("kraken", engine_keys)
        self.assertIn("opencv_rulings", engine_keys)
        self.assertIn("table_transformer", engine_keys)
        self.assertIn("yolo", engine_keys)
        self.assertIn("paddleocr", engine_keys)


if __name__ == "__main__":
    unittest.main()



